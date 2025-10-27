import os
import logging
from datetime import datetime, timedelta
from telegram import Update, InputFile
from telegram.ext import Application, CommandHandler, MessageHandler, filters, ContextTypes
import requests
from PyPDF2 import PdfReader
from docx import Document
from fpdf import FPDF

# === НАСТРОЙКИ ===
BOT_TOKEN = os.getenv("BOT_TOKEN")
CHANNEL_ID = int(os.getenv("CHANNEL_ID"))
CHANNEL_USERNAME = os.getenv("CHANNEL_USERNAME")

# Хранилище лимитов
user_limits = {}

logging.basicConfig(
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    level=logging.INFO
)
logger = logging.getLogger(__name__)

# Проверка подписки
async def is_subscribed(user_id: int) -> bool:
    url = f"https://api.telegram.org/bot{BOT_TOKEN}/getChatMember"
    params = {"chat_id": CHANNEL_ID, "user_id": user_id}
    try:
        response = requests.get(url, params=params, timeout=5)
        data = response.json()
        status = data.get("result", {}).get("status", "")
        return status in ["member", "administrator", "creator"]
    except Exception as e:
        logger.error(f"Ошибка проверки: {e}")
        return False

def get_daily_reset():
    now = datetime.utcnow()
    return now.replace(hour=0, minute=0, second=0, microsecond=0) + timedelta(days=1)

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    if not await is_subscribed(user.id):
        await update.message.reply_text(
            f"🔒 Подпишись на канал, чтобы использовать бота:\n"
            f"https://t.me/{CHANNEL_USERNAME}"
        )
        return
    await update.message.reply_text(
        "✨ Поддерживаю:\n"
        "• PDF → TXT\n"
        "• DOCX → TXT\n"
        "• TXT → PDF\n\n"
        "Отправь файл для конвертации!\n"
        "Лимит: 10 конвертаций в день."
    )

def check_limit(user_id: int) -> bool:
    now = datetime.utcnow()
    if user_id not in user_limits:
        user_limits[user_id] = {"count": 0, "reset_time": get_daily_reset()}
    user_data = user_limits[user_id]
    if now >= user_data["reset_time"]:
        user_data["count"] = 0
        user_data["reset_time"] = get_daily_reset()
    return user_data["count"] < 10

def increment_limit(user_id: int):
    if user_id in user_limits:
        user_limits[user_id]["count"] += 1

# Обработка файлов
async def handle_file(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    if not await is_subscribed(user.id):
        await update.message.reply_text(
            f"Подпишись на канал: https://t.me/{CHANNEL_USERNAME}"
        )
        return

    if not check_limit(user.id):
        await update.message.reply_text(
            "🚫 Достигнут лимит: 10 конвертаций в день.\n"
            "Завтра будет новый лимит!\n\n"
            f"Следи за обновлениями: https://t.me/{CHANNEL_USERNAME}"
        )
        return

    file = None
    if update.message.document:
        file = update.message.document
        mime_type = file.mime_type or ""
    else:
        await update.message.reply_text("Отправь файл (документ).")
        return

    if not file:
        return

    try:
        # Скачиваем
        file_obj = await context.bot.get_file(file.file_id)
        file_path = f"/tmp/temp_{user.id}_{file.file_unique_id}"
        await file_obj.download_to_drive(file_path)

        output_path = None
        caption = ""

        # PDF → TXT
        if file_path.lower().endswith(".pdf"):
            output_path = file_path.replace(".pdf", ".txt")
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n\n"
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(text)
            caption = "✅ PDF → TXT"

        # DOCX → TXT
        elif file_path.lower().endswith(".docx"):
            output_path = file_path.replace(".docx", ".txt")
            doc = Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs])
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(text)
            caption = "✅ DOCX → TXT"

        # TXT → PDF
        elif file_path.lower().endswith(".txt"):
            output_path = file_path.replace(".txt", ".pdf")
            pdf = FPDF()
            pdf.add_page()
            pdf.set_auto_page_break(auto=True, margin=15)
            pdf.set_font("Arial", size=12)
            with open(file_path, "r", encoding="utf-8") as f:
                for line in f:
                    pdf.cell(0, 10, txt=line.encode('latin-1', 'replace').decode('latin-1'), ln=True)
            pdf.output(output_path)
            caption = "✅ TXT → PDF"

        else:
            await update.message.reply_text(
                "❌ Поддерживаю только PDF, DOCX, TXT."
            )
            os.remove(file_path)
            return

        # Отправляем
        with open(output_path, "rb") as f:
            await update.message.reply_document(document=InputFile(f), caption=caption)

        # Убираем временные файлы
        os.remove(file_path)
        os.remove(output_path)
        increment_limit(user.id)

    except Exception as e:
        logger.error(f"Ошибка: {e}")
        await update.message.reply_text("❌ Ошибка конвертации. Попробуй другой файл.")
        if 'file_path' in locals() and os.path.exists(file_path):
            os.remove(file_path)
        if 'output_path' in locals() and os.path.exists(output_path):
            os.remove(output_path)

# Запуск
def main():
    application = Application.builder().token(BOT_TOKEN).build()
    application.add_handler(CommandHandler("start", start))
    application.add_handler(MessageHandler(filters.Document.ALL, handle_file))
    logger.info("Бот запущен!")
    application.run_polling()

if __name__ == "__main__":
    main()
